"""Documentazione del progetto Mirwada — rigenera Documentazione_Mirwada.docx.

Come si usa: aggiornare i testi in SEZIONI (e AGGIORNATO_AL), poi rilanciare:

    python docs/documentazione.py

Lo script rigenera sempre il docx da zero (idempotente). Le righe che iniziano
con "- " diventano elenchi puntati; le righe vuote separano i paragrafi.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

PROGETTO = "Mirwada"
AGGIORNATO_AL = "29.10.2026"  # data ultimo aggiornamento, formato GG.MM.AAAA

# Le 5 sezioni fisse della documentazione. Testo discorsivo, righe "- " = bullet.
SEZIONI = {
    "Come funziona": """Mirwada e' un action-RPG 2D top-down in pixel art (Godot 4.3, GDScript) con progressione a Pathway e Sequenze: 10 Pathway attivi in 4 gruppi completi, 10 Sequenze ciascuno dalla 9 alla 0, piu' i Pathway non-standard che avanzano per Boon. Il gioco si avvia con il libro piccolo e il panorama che scorre, si crea il personaggio (nome, Pathway, due talenti innati, modalita' facile o difficile) e si gioca in un mondo unico di 1200 x 1000 caselle.

Il principio architetturale e' uno: il codice implementa primitive parametriche (28 attive), i dati JSON le compongono in abilita'. I registri sono CHIUSI (primitive, 12 eventi tracciabili, tag, tipi di pagina e di gate): aggiungere una voce e' una decisione, non un gesto.

- Mondo: generato una volta dal generatore forge3d (seed 6) e salvato come dati in assets/world/ (tessere gia' composte, collisioni, regioni, oggetti di scena, edifici, il sottosuolo come secondo livello). scripts/world/mondo.gd costruisce solo i blocchi attorno al giocatore. Di notte il mondo si scurisce (luce.json); la fauna d'atmosfera scappa quando ci si avvicina.
- Citta' e villaggi: le case del generatore hanno le porte verso 18 interni disegnati e arredati (botteghe, case, torri, le quattro basi). Negli interni la camera inquadra la stanza sotto l'HUD.
- Spostamenti: a piedi, in barca dal porto, in volo con la piuma dei venti, dalla soglia del faro verso la Frontiera; le Marche si raggiungono solo volando o dal sottosuolo. I varchi delle sotto-aree sono ingressi verso un interno con il contenuto della zona.
- Contenuto: NPC con aspetto e ritratto, nemici e boss dei layout portati sul mondo, forzieri con tabelle di bottino, oggetti nelle case; ogni oggetto fisso si prende una volta sola (nel salvataggio).
- Base del giocatore: quattro basi (casa al porto, capanna nella Valle, rifugio nelle Marche, nave al cantiere navale) con tre stanze che si arredano col livello; la sede si sposta; fuori ci sono l'orto con le piante che crescono e le strutture costruite dalla pagina Base.
- Arte: personaggi a 64 px con tutte le animazioni (anche gesti, volo, barca e due forme di Enel), VFX di tutte le primitive colorati dalla palette del Pathway, icone di oggetti, stati e primitive, ritratti, panorama all'avvio.
- Il libro e' tutta l'interfaccia (scaffale, frontespizio, diagramma, inventario con la Base, mappa, dialoghi, negozio, colophon), sulle due facciate e senza scrivere sulla piega.
- data/: pathways, abilities, synergies, schema/ (registri), animations.json, vfx.json, audio.json, balance.json, world/ (regioni, layout, interni, bottino, fauna, mezzi, luce, cannocchiali), base/, structures/, ui/.
- tools/: validate_data.py (esce 0/1), run_tests.sh (suite + controllo degli errori nascosti), art/ (generatori rilanciabili).
- tests/: 1076 test headless.""",
    "Perché funziona così": """Tutto e' data-driven perche' il progetto e' enorme per una persona: aggiungere contenuto deve costare una riga di JSON e mai una classe. La prova: dieci Pathway completi con zero righe di codice che ne nominino uno.

- I nomi (opera di riferimento protetta) vivono solo nei dati e nelle chiavi i18n: il flag SERIAL_NUMBERS_FILED_OFF permette la rinominazione completa.
- L'arte la genera Claude con sprite-forge (forge3d): ogni asset esce da uno script rilanciabile con un seed, in palette Endesga 32, e si guarda ingrandito prima di consegnarlo. Mai PNG ritoccati a mano.
- Il mondo si genera una volta e si salva come dati: il gioco non rigenera nulla a runtime, e le correzioni si fanno nel generatore o nei dati (per esempio una casa in piu' per un interno nuovo di un villaggio la mette il generatore).
- Le decisioni di gioco stanno nei dati: come si arriva in una regione, dove stanno la barca e la nave, i colori della notte, i posti fuori dalla base, i costi delle strutture.
- La suite fallisce anche se un test si interrompe a meta' per un errore di script (tools/run_tests.sh): GDScript non ha eccezioni e prima quei test risultavano superati.
- Ogni story di gioco si guarda a schermo (screenshot sotto Xvfb) prima del commit: molti errori (un orto sotto un albero, una passerella irraggiungibile) li hanno trovati solo le schermate.""",
    "Come deployarlo": """Non c'e' ancora una build distribuibile: il gioco si avvia dall'editor o da riga di comando.

- Godot 4.3. Prima volta e dopo aver cambiato le PNG: godot --headless --path . --import.
- Python 3 per i dati (nessuna dipendenza); numpy e Pillow per tools/art/ e la mappa; python-docx per questa documentazione.
- Verifica completa: python tools/validate_data.py (esce 0) e bash tools/run_tests.sh (esce 0, GODOT=percorso se il binario non e' sul PATH).
- CI: GitHub Actions su push a main e sulle pull request (validator + suite con il controllo degli errori nascosti). Dal 29.10 tutto il lavoro e' su main; i controlli della CI si guardano solo quando li chiede l'utente.
- Rigenerare l'arte: tools/art/build_*.py; il mondo con tools/art/esporta_mondo.py (circa 1 ora per tutto; --edifici, --barca, --varchi, --sottosuolo, --solo-blocchi per le parti).
- Repo: https://github.com/YasserOudabashi/mirwada. Nota IP nel CLAUDE.md: non distribuibile ne' monetizzabile coi nomi attuali.""",
    "Cosa migliorare": """Dal controllo completo del 29.10.2026 (tutto segnato in prd.json):

- La storia dell'Atto II e III manca: 4 quest in tutto il gioco, le tribolazioni hanno solo contatori, e la prova finale (trib_1_0) non ha nessuna fonte che la segni superata: oggi il gioco non si puo' finire.
- Gli effetti delle primitive che colpiscono un altro stanno su chi lancia; alcuni durano troppo poco per vedersi.
- Il negozio e la pagina Indosso non mostrano le icone; le icone hanno una serie sola anche nella modalita' cupa; i sigilli incastonati non si vedono.
- La barca va alla velocita' a piedi; i pontili stretti del porto non si raggiungono a piedi; nelle stanze piccole i nemici arrivano subito addosso.
- Mancano le luci di notte (finestre, lampioni, faro), i panorami per regione, le creature (evocazioni, antagonisti), nuoto e arrampicata, il cannocchiale da tasca.
- 7 percorsi di fusione su 8 sono stub; l'inglese manca per 1496 testi su 1784.
- Da provare giocando: i tempi d'attacco nuovi, la durata dei tragitti, il caricamento del mondo su Windows.""",
    "Prossimi punti": """In ordine di priorita' (prd.json, 57 story aperte):

- Miglioramenti rapidi: effetti sul bersaglio, icone nel negozio e in Indosso, barca piu' veloce, pontili, cannocchiale, gesti, nemici nelle stanze, sigilli (US-8B44..US-8B73e).
- Arte rimasta: cannocchiale da tasca, luci di notte, panorami per regione, evocazioni e antagonisti, nuoto, arrampicata (US-8B26..US-8B92).
- Fase 13 Atto II e III (US-1301..US-1313): quello che serve per arrivare alla fine.
- Fase 12 Mondo vivo: prima riscrivere il PRD sul mondo generato (US-1200), poi US-1201..US-1216.
- Fusioni (US-7B01..07) e traduzione inglese (US-EN01..05).""",
}

BLU = RGBColor(0x1A, 0x23, 0x7E)
OUTPUT = Path(__file__).parent / f"Documentazione_{PROGETTO}.docx"


def _p(doc, testo, *, bold=False, size=None, color=None):
    par = doc.add_paragraph()
    run = par.add_run(testo)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return par


def _corpo(doc, testo):
    """Rende un blocco di testo: bullet per righe '- ', paragrafi per il resto."""
    for riga in [r.strip() for r in testo.strip().splitlines()]:
        if not riga:
            continue
        if riga.startswith("- "):
            doc.add_paragraph(riga[2:], style="List Bullet")
        else:
            doc.add_paragraph(riga)


def genera():
    doc = Document()
    _p(doc, f"Documentazione — {PROGETTO}", bold=True, size=16, color=BLU)
    _p(doc, f"Aggiornata al: {AGGIORNATO_AL}", size=10)
    for titolo, testo in SEZIONI.items():
        _p(doc, titolo, bold=True, size=13, color=BLU)
        if testo.strip():
            _corpo(doc, testo)
        else:
            _p(doc, "(da compilare)")
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    percorso = genera()
    assert percorso.exists() and percorso.stat().st_size > 0, "docx non generato"
    print(f"OK — generato {percorso}")
